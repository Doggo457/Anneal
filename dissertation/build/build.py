"""Fill the Stirling dissertation outline template with content.py."""
import copy, datetime, io, re, struct, sys, importlib.util
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = sys.argv[1]
TEMPLATE, OUT = sys.argv[2], sys.argv[3]
spec = importlib.util.spec_from_file_location("content", f"{HERE}/content.py")
C = importlib.util.module_from_spec(spec); spec.loader.exec_module(C)

def clean_png(path):
    """Figure bytes with metadata chunks removed: only image data and colour/size chunks are kept."""
    data = open(path, "rb").read()
    if data[:8] != b"\x89PNG\r\n\x1a\n":
        return path
    keep = {b"IHDR", b"PLTE", b"IDAT", b"IEND", b"tRNS", b"gAMA", b"cHRM", b"sRGB", b"iCCP", b"pHYs", b"sBIT"}
    out, i = [data[:8]], 8
    while i < len(data):
        n = struct.unpack(">I", data[i:i + 4])[0]
        if data[i + 4:i + 8] in keep:
            out.append(data[i:i + 12 + n])
        i += 12 + n
    return io.BytesIO(b"".join(out))

d = docx.Document(TEMPLATE)
body = d.element.body

# ---------- citation numbering by first appearance ----------
order = []
def scan(s):
    for k in re.findall(r"\[@([\w-]+)\]", s):
        if k not in order: order.append(k)
def walk(x):
    if isinstance(x, str): scan(x)
    elif isinstance(x, (list, tuple)):
        for y in x: walk(y)
walk(C.BLOCKS)
missing = [k for k in order if k not in C.REFS]
if missing: sys.exit(f"missing refs: {missing}")
unused = [k for k in C.REFS if k not in order]
if unused: print("WARNING unused refs:", unused)
num = {k: i + 1 for i, k in enumerate(order)}
leftover = re.findall(r"<<\w+>>", repr(C.BLOCKS) + repr(C.REFS))
if leftover: sys.exit(f"unfilled tokens: {leftover}")

# ---------- inline runs ----------
TOK = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\[@[\w-]+\]|\{\{.+?\}\}|`.+?`)")

def mk_run(text, bold=False, italic=False, highlight=False, mono=False, size=None, color=None):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if mono:
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs"): f.set(qn(a), "Consolas")
        rpr.append(f)
    if bold: rpr.append(OxmlElement("w:b"))
    if italic: rpr.append(OxmlElement("w:i"))
    if color:
        c = OxmlElement("w:color"); c.set(qn("w:val"), color); rpr.append(c)
    if size:
        s = OxmlElement("w:sz"); s.set(qn("w:val"), str(size)); rpr.append(s)
    if highlight:
        h = OxmlElement("w:highlight"); h.set(qn("w:val"), "yellow"); rpr.append(h)
    if color == "0563C1":
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    if len(rpr): r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t)
    return r

def cite(key, size=None):
    f = OxmlElement("w:fldSimple")
    f.set(qn("w:instr"), f" REF _RefAnneal{num[key]} \\r \\h ")
    f.append(mk_run(f"[{num[key]}]", size=size))
    return f

def fill(p_el, text, size=None, base_bold=False):
    for part in TOK.split(text):
        if not part: continue
        if part.startswith("**"): p_el.append(mk_run(part[2:-2], bold=True, size=size))
        elif part.startswith("[@"): p_el.append(cite(part[2:-1], size))
        elif part.startswith("{{"): p_el.append(mk_run("[" + part[2:-2] + "]", highlight=True, size=size, bold=base_bold))
        elif part.startswith("`"): p_el.append(mk_run(part[1:-1], mono=True, size=size or 19))
        elif part.startswith("*") and len(part) > 2: p_el.append(mk_run(part[1:-1], italic=True, size=size, bold=base_bold))
        else: p_el.append(mk_run(part, size=size, bold=base_bold))

def para(style, text="", **kw):
    p = d.add_paragraph(style=style)
    if text: fill(p._p, text, **kw)
    return p

def set_ppr(p, jc=None, after=None, before=None, ind_left=None, hanging=None, keep_next=False):
    ppr = p._p.get_or_add_pPr()
    if keep_next: ppr.append(OxmlElement("w:keepNext"))
    if after is not None or before is not None:
        sp = OxmlElement("w:spacing")
        if before is not None: sp.set(qn("w:before"), str(before))
        if after is not None: sp.set(qn("w:after"), str(after))
        ppr.append(sp)
    if ind_left is not None:
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), str(ind_left))
        if hanging: ind.set(qn("w:hanging"), str(hanging))
        ppr.append(ind)
    if jc:
        j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); ppr.append(j)

# ---------- cover page ----------
cover = [p for p in d.paragraphs[:8]]
def set_text(p, text):
    rpr = None
    if p.runs and p.runs[0]._r.rPr is not None: rpr = copy.deepcopy(p.runs[0]._r.rPr)
    for r in list(p._p):
        if r.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:fldSimple")): p._p.remove(r)
    start = len(p._p)
    fill(p._p, text)
    if rpr is not None:
        for r in p._p[start:]:
            if r.tag == qn("w:r"):
                old = r.find(qn("w:rPr")); new = copy.deepcopy(rpr)
                if old is not None:
                    for ch in old:  # keep highlight/bold from markup
                        new.append(ch)
                    r.remove(old)
                r.insert(0, new)
by_text = {p.text.strip(): p for p in cover}
set_text(by_text["Project Title"], C.TITLE)
set_text(by_text["Author Name"], C.AUTHOR)
set_text(by_text["Author Student ID"], C.STUDENT_ID)
set_text(by_text["Supervisor"], C.SUPERVISOR)
set_text(by_text["October 202X"], C.DATE)

# ---------- remove template body after the cover section ----------
kids = list(body.iterchildren())
start = next(i for i, el in enumerate(kids) if el.tag == qn("w:p") and el.find(".//" + qn("w:sectPr")) is not None) + 1
for el in kids[start:]:
    if el.tag != qn("w:sectPr"): body.remove(el)
sectPr = body.find(qn("w:sectPr"))

# ---------- tables ----------
def add_table(widths, header, rows, size=20, first_col_bold=True):
    tbl = OxmlElement("w:tbl")
    tp = OxmlElement("w:tblPr")
    st = OxmlElement("w:tblStyle"); st.set(qn("w:val"), "GridTable4-Accent31"); tp.append(st)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(sum(widths))); tw.set(qn("w:type"), "dxa"); tp.append(tw)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tp.append(lay)
    look = OxmlElement("w:tblLook")
    for k, v in (("w:val", "04A0"), ("w:firstRow", "1"), ("w:lastRow", "0"), ("w:firstColumn", "1" if first_col_bold else "0"),
                 ("w:lastColumn", "0"), ("w:noHBand", "0"), ("w:noVBand", "1")):
        look.set(qn(k), v)
    tp.append(look); tbl.append(tp)
    g = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(w)); g.append(gc)
    tbl.append(g)
    for ri, row in enumerate([header] + rows):
        tr = OxmlElement("w:tr")
        trpr = OxmlElement("w:trPr"); trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0: trpr.append(OxmlElement("w:tblHeader"))
        tr.append(trpr)
        for w, cell in zip(widths, row):
            tc = OxmlElement("w:tc"); tcp = OxmlElement("w:tcPr")
            cw = OxmlElement("w:tcW"); cw.set(qn("w:w"), str(w)); cw.set(qn("w:type"), "dxa"); tcp.append(cw); tc.append(tcp)
            p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
            ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), "BodyText"); ppr.append(ps)
            if len(rows) <= 10 and ri < len(rows): ppr.append(OxmlElement("w:keepNext"))  # keep short tables on one page
            sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "20"); sp.set(qn("w:after"), "20"); ppr.append(sp)
            j = OxmlElement("w:jc"); j.set(qn("w:val"), "left"); ppr.append(j)
            p.append(ppr); fill(p, cell, size=size); tc.append(p); tr.append(tc)
        tbl.append(tr)
    sectPr.addprevious(tbl)
    return tbl

def caption(text):
    p = para("Body Centre")
    p._p.append(mk_run(text.split(".")[0] + ".", bold=True, size=20))
    rest = text[len(text.split(".")[0]) + 1:]
    if rest: p._p.append(mk_run(rest, size=20))
    set_ppr(p, before=120, after=60, keep_next=True)

# ---------- body ----------
plan_hdr = ["Week", "Project Task", "Deliverables"]
for b in C.BLOCKS:
    kind = b[0]
    if kind == "h1": para("Heading 1", b[1])
    elif kind == "h2": para("Heading 2", b[1])
    elif kind == "h3": para("Heading 3", b[1])
    elif kind == "bf": para("Body First", b[1])
    elif kind == "bt": para("Body Text", b[1])
    elif kind == "obj":
        p = para("Body Text")
        p._p.append(mk_run(b[1], bold=True))
        tab = OxmlElement("w:r"); tab.append(OxmlElement("w:tab")); p._p.append(tab)
        fill(p._p, b[2])
        set_ppr(p, ind_left=567, hanging=567, after=60)
        ppr = p._p.pPr
        tabs = OxmlElement("w:tabs"); t = OxmlElement("w:tab"); t.set(qn("w:val"), "left"); t.set(qn("w:pos"), "567"); tabs.append(t)
        ppr.insert(1, tabs)
    elif kind == "fig":
        p = para("Body Centre"); set_ppr(p, keep_next=True, before=120)
        p.add_run().add_picture(clean_png(f"{HERE}/{b[1]}"), width=Inches(5.7))
        para("Figure", b[2])
    elif kind == "table":
        caption(b[1]); add_table(b[2], b[3], b[4])
        para("Body Text").paragraph_format.space_after = 0
    elif kind == "plan":
        add_table([1700, 4200, 2396], plan_hdr, b[2], size=20)
        para("Body Text").paragraph_format.space_after = 0
    elif kind == "code":
        lines = b[2]
        for i, line in enumerate(lines):
            p = para("Body Text")
            p._p.append(mk_run(line.replace(" ", " "), mono=True, size=17))
            set_ppr(p, jc="left", after=0, before=120 if i == 0 else 0, ind_left=284, keep_next=True)
        caption(b[1])
    else:
        sys.exit(f"unknown block {kind}")

# ---------- references ----------
# widen the Reference list's number column so [10]+ keeps a gap (Word and LibreOffice)
numxml = d.part.numbering_part.element
ref_style = d.styles["Reference"].element
ref_numid = ref_style.find(".//" + qn("w:numId")).get(qn("w:val"))
for n in numxml.findall(qn("w:num")):
    if n.get(qn("w:numId")) == ref_numid:
        absid = n.find(qn("w:abstractNumId")).get(qn("w:val"))
for a in numxml.findall(qn("w:abstractNum")):
    if a.get(qn("w:abstractNumId")) == absid:
        lvl = next(l for l in a.findall(qn("w:lvl")) if l.get(qn("w:ilvl")) == "0")
        lppr = lvl.find(qn("w:pPr"))
        for ch in list(lppr): lppr.remove(ch)
        tabs = OxmlElement("w:tabs"); t = OxmlElement("w:tab"); t.set(qn("w:val"), "num"); t.set(qn("w:pos"), "567"); tabs.append(t)
        ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "567"); ind.set(qn("w:hanging"), "567")
        lppr.append(tabs); lppr.append(ind)
para("Unnumbered 1", "References")
for i, k in enumerate(order, 1):
    text, url, *suffix = C.REFS[k]
    p = para("Reference")
    bm = OxmlElement("w:bookmarkStart"); bm.set(qn("w:id"), str(900 + i)); bm.set(qn("w:name"), f"_RefAnneal{i}")
    p._p.insert(1 if p._p.pPr is not None else 0, bm)
    fill(p._p, text + " ")
    p._p.append(mk_run("[Online]. Available: "))
    rid = d.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    h.append(mk_run(url, color="0563C1")); p._p.append(h)
    if suffix: p._p.append(mk_run(suffix[0]))
    be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), str(900 + i)); p._p.append(be)
    ppr = p._p.pPr
    tabs = OxmlElement("w:tabs")
    for v, pos in (("clear", "426"), ("num", "567")):
        t = OxmlElement("w:tab"); t.set(qn("w:val"), v); t.set(qn("w:pos"), pos); tabs.append(t)
    ppr.insert(2 if ppr.find(qn("w:numPr")) is not None else 1, tabs)
    set_ppr(p, ind_left=567, hanging=567, jc="left")

d.core_properties.title = C.TITLE
author = re.sub(r"\{\{.*?\}\}", "", C.AUTHOR).strip()
d.core_properties.author = author
d.core_properties.last_modified_by = author
d.core_properties.modified = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
d.save(OUT)
print("saved", OUT, "refs:", {k: num[k] for k in order})
